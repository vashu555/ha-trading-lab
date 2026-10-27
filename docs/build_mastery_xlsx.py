"""Build the HA Trading Lab mastery checklist as a plain Excel sheet (one row per task).

    python docs/build_mastery_xlsx.py      → docs/HA_Lab_Mastery_Checklist.xlsx

The content comes from build_mastery_pdf.py, so the two never disagree. Needs openpyxl.
"""

from __future__ import annotations

import html
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import build_mastery_pdf as m

OUT = Path(__file__).resolve().parent / "HA_Lab_Mastery_Checklist.xlsx"
FONT = "Arial"


def plain(s: str) -> str:
    return html.unescape(re.sub(r"<[^>]+>", "", s))


def rows():
    """(checkpoint, component, week, level, task, command / files, expected)"""
    for x in m.START:
        yield 0, "Own the lab", 1, "Setup", plain(x), "", ""
    yield 0, "Own the lab", 1, "Pass", m.START_PASS, "", ""
    for cp in m.CHECKPOINTS:
        n, t, w = cp["n"], cp["title"], cp["week"]
        yield n, t, w, "1 Explain", "Answer in your own words:\n" + "\n".join("- " + plain(q) for q in cp["explain"]), "", ""
        yield n, t, w, "2 Inspect", "Run and read the output", "\n".join(cp["inspect"]), ""
        for cmd, expect in cp["brk"]:
            yield n, t, w, "3 Break", "Write your prediction first, then run", cmd, expect
        yield n, t, w, "4 Read", "Explain each file / function", "\n".join(cp["read"]), ""
        yield n, t, w, "5 Change", plain(cp["change"]), "", "Audit stays 0; revert with git"
    for x in m.FINAL:
        yield "Final", "Guarantee", 4, "Final", plain(x), "", ""


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"
    head = ["#", "Checkpoint", "Component", "Week", "Level", "Task", "Command / files", "Expected result",
            "Done", "Date", "Checked by", "Notes"]
    widths = [5, 11, 24, 7, 11, 48, 58, 34, 8, 12, 14, 24]
    ws.append(head)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + c)].width = w
        cell = ws.cell(1, c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    n = 0
    for r in rows():
        n += 1
        ws.append([n, *r[:2], r[2], *r[3:], None, None, None, None])
    for row in ws.iter_rows(min_row=2, max_row=n + 1):
        for cell in row:
            cell.font = Font(name=FONT, size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[6].font = Font(name="Consolas", size=9)              # commands
        row[9].number_format = "dd-mmm-yyyy"
    done = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(done)
    done.add(f"I2:I{n + 1}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{n + 1}"

    # progress per checkpoint: counts the "Yes" in the Done column
    pr = wb.create_sheet("Progress")
    pr.append(["Checkpoint", "Component", "Tasks", "Done", "% done"])
    pr.append(["How to use: fill Done (Yes/No), Date and Checked by on the Checklist sheet; this sheet updates itself.",
               None, None, None, None])
    pr.merge_cells("A2:E2")
    pr["A2"].font = Font(name=FONT, size=9, italic=True, color="595959")
    cps = [(0, "Own the lab")] + [(cp["n"], cp["title"]) for cp in m.CHECKPOINTS] + [("Final", "Guarantee")]
    for i, (k, t) in enumerate(cps, 3):
        key = f'"{k}"' if isinstance(k, str) else k
        pr.append([k, t, f"=COUNTIF(Checklist!$B:$B,{key})", f'=COUNTIFS(Checklist!$B:$B,{key},Checklist!$I:$I,"Yes")',
                   f"=IF(C{i}=0,0,D{i}/C{i})"])
        pr.cell(i, 5).number_format = "0%"
    last = 2 + len(cps)
    pr.append(["Total", None, f"=SUM(C3:C{last})", f"=SUM(D3:D{last})", f"=IF(C{last + 1}=0,0,D{last + 1}/C{last + 1})"])
    pr.cell(last + 1, 5).number_format = "0%"
    for c, w in zip("ABCDE", [12, 30, 8, 8, 9]):
        pr.column_dimensions[c].width = w
    for row in pr.iter_rows():
        for cell in row:
            if cell.row != 2:
                cell.font = Font(name=FONT, size=10, bold=cell.row in (1, last + 1))
    for c in range(1, 6):
        pr.cell(1, c).fill = PatternFill("solid", fgColor="1F3864")
        pr.cell(1, c).font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f"wrote {OUT.name} ({n} tasks)")


if __name__ == "__main__":
    build()
